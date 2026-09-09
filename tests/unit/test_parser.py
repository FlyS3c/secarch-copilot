from __future__ import annotations

from pathlib import Path

from docx import Document

from app.ingestion.parser import parse_document


def test_markdown_preserves_heading_provenance(tmp_path: Path) -> None:
    path = tmp_path / "fixture.md"
    path.write_text(
        "# Identity\nUse phishing-resistant MFA.\n\n## Logging\nCentralize logs.",
        encoding="utf-8",
    )

    sections = parse_document(path, "markdown")
    assert sections[0].location.startswith("Identity")
    assert "phishing-resistant MFA" in sections[0].text
    assert sections[1].location.startswith("Logging")


def test_docx_preserves_heading_provenance(tmp_path: Path) -> None:
    path = tmp_path / "fixture.docx"
    document = Document()
    document.add_heading("Trust Boundaries", level=1)
    document.add_paragraph("Validate data before it crosses the boundary.")
    document.save(path)

    sections = parse_document(path, "docx")
    assert sections[0].location == "Trust Boundaries"
    assert "Validate data" in sections[0].text


def test_empty_text_file_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")

    try:
        parse_document(path, "txt")
    except ValueError as exc:
        assert "No extractable text" in str(exc)
    else:
        raise AssertionError("Expected empty document to be rejected")
